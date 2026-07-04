"""
json_to_docx.py

Converts Whisper-style transcription JSON files into .docx Word documents.

Each input JSON is expected to have (at least) these keys:
    - "source_path"     : original audio file path (used to build a title)
    - "duration_s"       : audio duration in seconds
    - "detected_language" : language code detected by Whisper
    - "corrected_text"    : the cleaned/corrected transcript text (this is
                            what gets written into the docx)

Usage:
    python3 json_to_docx.py
(Just edit SRC_DIR / FILES below to point at your own JSON files, or pass
file paths as command-line arguments.)
"""

import json
import os
import re
import sys

from docx import Document
from docx.shared import Pt

# Folder that contains the input JSON files and where the .docx files
# will be written.
SRC_DIR = "/home/mpinganzima/Downloads/Ibitero_Kigeli"

# List the JSON files to convert. If you pass file paths on the command
# line instead (python3 json_to_docx.py file1.json file2.json ...),
# those will be used instead of this list.
FILES = [
    "Ibitero_1__KIGELI_IV_Rwabugiri_by_Mwanafunzi_na_MUNANA__rjj4w6_qNTM_.json",
    "Ibitero_2__KIGELI_IV_Rwabugiri_by_Mwanafunzi_na_MUNANA__Al38IV64Oc8_.json",
    "Ibitero_3__KIGELI_IV_Rwabugiri_by_Mwanafunzi_na_MUNANA__nc3re_iPssY_.json",
    "Ibitero_4__KIGELI_IV_Rwabugiri_by_Mwanafunzi_na_MUNANA___M8LRb423a0_.json",
]

# How many sentences to group into a single paragraph. The raw transcript
# text has no line breaks at all, so without this everything would end up
# as one giant wall-of-text paragraph.
SENTENCES_PER_PARAGRAPH = 5


def make_title(source_path, fallback_name):
    """
    Build a readable document title.

    Prefer the original audio filename (stored in "source_path" inside the
    JSON), since it usually has the real video/episode title. If that is
    missing, fall back to the JSON filename itself.
    """
    base = os.path.basename(source_path) if source_path else fallback_name
    base = os.path.splitext(base)[0]        # drop .mp3 / .json extension
    base = base.replace("：", ":")          # normalize fullwidth colon
    return base


def split_into_sentences(text):
    """
    Split a block of text into a list of sentences.

    We split right after '.', '!' or '?' followed by whitespace. This is a
    simple heuristic (not perfect linguistics) but works fine for cleaning
    up transcript text that has no paragraph breaks.
    """
    parts = re.split(r'(?<=[.!?])\s+', text.strip())
    return [p.strip() for p in parts if p.strip()]


def group_into_paragraphs(sentences, per_paragraph=SENTENCES_PER_PARAGRAPH):
    """
    Group a flat list of sentences into paragraphs of `per_paragraph`
    sentences each, joined back into single strings.
    """
    paragraphs = []
    for i in range(0, len(sentences), per_paragraph):
        chunk = sentences[i:i + per_paragraph]
        paragraphs.append(" ".join(chunk))
    return paragraphs


def convert_json_to_docx(json_path):
    """
    Read one transcription JSON file and write a matching .docx file
    next to it (same name, .docx extension instead of .json).
    """
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    title = make_title(data.get("source_path"), os.path.basename(json_path))
    duration = data.get("duration_s", 0)
    language = data.get("detected_language", "")
    text = data["corrected_text"]  # use the cleaned/corrected transcript

    sentences = split_into_sentences(text)
    paragraphs = group_into_paragraphs(sentences)

    # --- Build the Word document ---
    doc = Document()

    # Base font for the whole document.
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(12)

    # Title heading.
    doc.add_heading(title, level=1)

    # Small italic metadata line (duration + detected language).
    meta_paragraph = doc.add_paragraph()
    meta_run = meta_paragraph.add_run(
        f"Duration: {duration:.0f}s | Language: {language}"
    )
    meta_run.italic = True
    meta_run.font.size = Pt(10)

    doc.add_paragraph()  # blank line for spacing

    # Body: one Word paragraph per group of sentences.
    for paragraph_text in paragraphs:
        doc.add_paragraph(paragraph_text)

    out_path = os.path.splitext(json_path)[0] + ".docx"
    doc.save(out_path)
    print(f"Wrote {out_path} ({len(paragraphs)} paragraphs, {len(sentences)} sentences)")
    return out_path


def main():
    # If file paths are given on the command line, use those.
    # Otherwise, fall back to the FILES list above (resolved against SRC_DIR).
    if len(sys.argv) > 1:
        json_paths = sys.argv[1:]
    else:
        json_paths = [os.path.join(SRC_DIR, name) for name in FILES]

    for json_path in json_paths:
        convert_json_to_docx(json_path)


if __name__ == "__main__":
    main()
